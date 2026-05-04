import os, asyncio, httpx, json, secrets, string, re, time
from datetime import datetime
from telegram import Update, ReplyKeyboardMarkup, KeyboardButton, InlineKeyboardButton, InlineKeyboardMarkup
from telegram.ext import ApplicationBuilder, CommandHandler, MessageHandler, CallbackQueryHandler, filters, ContextTypes
from flask import Flask
from threading import Thread
from docx import Document
from PIL import Image

# --- কনফিগারেশন ---
TOKEN = "8681613014:AAGhaxmwY5Xv_F1FDDyccUwIqT05WXUAQ3M"
GEMINI_KEY = "AIzaSyBLMcfkOSvW1s-6dEQ14p_T__on6ZmlXWY"
ADMIN_ID = 6365947320
CHANNELS = ["@shantirahmana", "@extramaill", "@allfileconverter"]

# --- ডাটাবেজ ---
DB_FILE = "database.json"
def load_db():
    if not os.path.exists(DB_FILE): return {"users": [], "vip": [], "stats": 0}
    with open(DB_FILE, "r") as f: return json.load(f)

def save_db(db):
    with open(DB_FILE, "w") as f: json.dump(db, f, indent=4)

# --- Force Join Check ---
async def is_subscribed(bot, user_id):
    for channel in CHANNELS:
        try:
            member = await bot.get_chat_member(chat_id=channel, user_id=user_id)
            if member.status in ['left', 'kicked']: return False
        except: return False # বট চ্যানেলে এডমিন না থাকলে এটি ফলস হবে
    return True

# --- কিবোর্ড মেনুসমূহ ---
def main_menu():
    return ReplyKeyboardMarkup([
        [KeyboardButton("🎬 Media Downloader"), KeyboardButton("📄 PDF & Office")],
        [KeyboardButton("📝 Resume Builder"), KeyboardButton("🖼 Image Tools")],
        [KeyboardButton("🤖 AI Chat Assistant"), KeyboardButton("🛠 Utility Tools")],
        [KeyboardButton("⚙️ Admin Panel"), KeyboardButton("👤 My Profile")]
    ], resize_keyboard=True)

# --- স্টার্ট কমান্ড ---
async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    db = load_db()
    if user_id not in db['users']:
        db['users'].append(user_id)
        save_db(db)

    # Force Join Check
    if not await is_subscribed(context.bot, user_id):
        btns = [[InlineKeyboardButton("📢 Join Channel 1", url="https://t.me/shantirahmana")],
                [InlineKeyboardButton("📢 Join Channel 2", url="https://t.me/extramaill")],
                [InlineKeyboardButton("📢 Join Channel 3", url="https://t.me/allfileconverter")],
                [InlineKeyboardButton("✅ I have Joined", callback_data="check_subs")]]
        await update.message.reply_text(
            "❌ **অ্যাক্সেস ডিনাইড!**\nবটটি ব্যবহার করতে আমাদের ৩টি চ্যানেলে জয়েন থাকা বাধ্যতামূলক।",
            reply_markup=InlineKeyboardMarkup(btns)
        )
        return

    await update.message.reply_text(
        f"🚀 **All File Converter-এ স্বাগতম!**\nআমি আপনার প্রফেশনাল মাল্টি-পারপাস বট।",
        reply_markup=main_menu()
    )

# --- এডমিন প্যানেল ---
async def admin_panel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID: return
    db = load_db()
    msg = f"⚙️ **Admin Panel**\n\n👥 মোট ইউজার: {len(db['users'])}\n🌟 ভিআইপি: {len(db['vip'])}\n\nব্রডকাস্ট করতে লিখুন: `/bc আপনার মেসেজ`"
    await update.message.reply_text(msg, parse_mode="Markdown")

async def broadcast(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if update.effective_user.id != ADMIN_ID: return
    msg = " ".join(context.args)
    db = load_db()
    count = 0
    for user in db['users']:
        try:
            await context.bot.send_message(chat_id=user, text=f"📢 **Admin Message:**\n\n{msg}")
            count += 1
        except: pass
    await update.message.reply_text(f"✅ {count} জন ইউজারকে মেসেজ পাঠানো হয়েছে।")

# --- Resume Builder (Step-by-Step) ---
async def resume_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['state'] = 'CV_NAME'
    await update.message.reply_text("📝 **Resume Builder শুরু হয়েছে!**\nআপনার পূর্ণ নাম লিখুন:")

async def handle_cv(update: Update, context: ContextTypes.DEFAULT_TYPE):
    state = context.user_data.get('state')
    text = update.message.text

    if state == 'CV_NAME':
        context.user_data['cv_name'] = text
        context.user_data['state'] = 'CV_SKILLS'
        await update.message.reply_text("✅ আপনার স্কিলগুলো লিখুন (যেমন: Python, Graphic Design):")
    
    elif state == 'CV_SKILLS':
        context.user_data['cv_skills'] = text
        context.user_data['state'] = 'CV_EXP'
        await update.message.reply_text("✅ আপনার কাজের অভিজ্ঞতা বা শিক্ষা সম্পর্কে লিখুন:")

    elif state == 'CV_EXP':
        # Resume তৈরি করা
        name = context.user_data['cv_name']
        skills = context.user_data['cv_skills']
        exp = text
        
        doc = Document()
        doc.add_heading(f'Resume of {name}', 0)
        doc.add_heading('Skills:', level=1)
        doc.add_paragraph(skills)
        doc.add_heading('Experience & Education:', level=1)
        doc.add_paragraph(exp)
        
        path = f"resume_{update.effective_user.id}.docx"
        doc.save(path)
        await update.message.reply_document(document=open(path, 'rb'), caption="✅ আপনার প্রফেশনাল রেজ্যুমে তৈরি!")
        os.remove(path)
        context.user_data.clear()

# --- মূল মেসেজ হ্যান্ডলার ---
async def handle_all(update: Update, context: ContextTypes.DEFAULT_TYPE):
    text = update.message.text
    user_id = update.effective_user.id

    # বাটন লজিক
    if text == "📝 Resume Builder":
        await resume_start(update, context)
        return
    
    elif text == "⚙️ Admin Panel":
        await admin_panel(update, context)
        return

    elif text == "🎬 Media Downloader":
        await update.message.reply_text("🎬 ভিডিওর লিংক দিন (YT/FB/IG/TikTok)।")
        return

    # Resume Builder এর স্টেট চেক
    if context.user_data.get('state', '').startswith('CV_'):
        await handle_cv(update, context)
        return

# --- Flask Server ---
app_web = Flask('')
@app_web.route('/')
def home(): return "All File Converter is Active!"
def keep_alive(): Thread(target=lambda: app_web.run(host='0.0.0.0', port=8080)).start()

if __name__ == '__main__':
    keep_alive()
    app = ApplicationBuilder().token(TOKEN).build()
    
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("bc", broadcast))
    app.add_handler(MessageHandler(filters.TEXT & (~filters.COMMAND), handle_all))
    
    print("🚀 All File Converter Bot Is LIVE!")
    app.run_polling()
